"""Seed workbook importer (FR-16).

Reads the "Top 100 AI agents by domain" workbook (12 domain sheets of 100
ranked rows, one Source Catalog sheet) and produces canonical resources plus
one domain-score row per sheet row. Every score is recomputed from its
components with ``config/scoring.json``; the import fails if any recomputed
score differs from the workbook by more than ``tolerance`` (default 0.1).

Requires ``openpyxl`` (``pip install agentdossier[connectors]``); the rest of
the core package stays standard-library only.
"""

from __future__ import annotations

import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..models import add_source, new_resource
from ..score import COMPONENTS, ScoringConfig, score_for_domain
from ..util import load_config, slug

SOURCE_SYSTEM = "seed_xlsx"

# Workbook sheet titles are truncated to 31 characters by Excel.
SHEET_TO_DOMAIN = {
    "Business & Enterprise Operation": "business",
    "Finance & Banking": "finance",
    "Healthcare & Life Sciences": "healthcare",
    "Regulatory, Risk & Compliance": "regulatory",
    "Technology & Software Engineeri": "technology",
    "Cybersecurity": "cybersecurity",
    "Sales, Marketing & Customer Exp": "sales",
    "Legal": "legal",
    "Data & Analytics": "data",
    "Supply Chain, Manufacturing & L": "supply_chain",
    "Insurance": "insurance",
    "Government & Public Sector": "government",
}

HEADER = [
    "Rank",
    "Agent / Product",
    "Category",
    "Vendor",
    "Ecosystem / Deployment",
    "Project / Product URL",
    "Observed popularity/status",
    "License / Commercial",
    "Adoption",
    "Trust",
    "Health",
    "Ecosystem",
    "Domain fit",
    "Governance / compliance fit",
    "Docs",
    "Domain score",
    "Why included",
    "Primary source",
    "Last verified",
]

_TYPE_RULES = (
    (re.compile(r"marketplace|discovery hub|community hub", re.I), "marketplace"),
    (re.compile(r"\bsdk\b|framework|orchestration|runtime|toolkit", re.I), "framework"),
    (re.compile(r"platform", re.I), "platform"),
    (re.compile(r"\bmcp\b|server", re.I), "mcp_server"),
    (re.compile(r"\btool\b|guardrail", re.I), "tool"),
)


def infer_resource_type(category: str | None) -> str:
    for pattern, rtype in _TYPE_RULES:
        if category and pattern.search(category):
            return rtype
    return "agent"


@dataclass
class SeedRow:
    domain: str
    rank: int
    name: str
    category: str | None
    vendor: str | None
    ecosystem: str | None
    url: str | None
    popularity: str | None
    license: str | None
    components: dict[str, float | None]
    workbook_score: float
    why: str | None
    primary_source: str | None
    last_verified: str | None


@dataclass
class SeedImport:
    snapshot_date: str
    resources: dict[str, dict[str, Any]]
    rows: list[SeedRow]
    max_score_error: float
    errors: list[str] = field(default_factory=list)
    source_catalog: list[dict[str, str]] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value if value and value.upper() != "N/A" else None
    return value


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_workbook(path: str | Path) -> tuple[str, list[SeedRow], list[dict[str, str]]]:
    try:
        import openpyxl  # noqa: PLC0415 - optional dependency
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "openpyxl is required to read the seed workbook: pip install 'agentdossier[connectors]'"
        ) from exc

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl warns about unsupported conditional-format extensions
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return _read_sheets(wb)


def _read_sheets(wb: Any) -> tuple[str, list[SeedRow], list[dict[str, str]]]:

    snapshot_date = "unknown"
    if "Methodology" in wb.sheetnames:
        for row in wb["Methodology"].iter_rows(values_only=True):
            if row and row[0] == "Snapshot date" and row[1]:
                snapshot_date = str(row[1])[:10]

    source_catalog: list[dict[str, str]] = []
    if "Source Catalog" in wb.sheetnames:
        rows = list(wb["Source Catalog"].iter_rows(values_only=True))
        for r in rows[1:]:
            if r and r[0]:
                source_catalog.append(
                    {
                        "url": str(r[0]),
                        "product": str(r[1] or ""),
                        "vendor": str(r[2] or ""),
                        "domains": str(r[3] or ""),
                    }
                )

    rows_out: list[SeedRow] = []
    for title, domain in SHEET_TO_DOMAIN.items():
        if title not in wb.sheetnames:
            raise ValueError(f"seed workbook is missing sheet '{title}'")
        ws = wb[title]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        if header[: len(HEADER)] != HEADER:
            raise ValueError(f"sheet '{title}' has an unexpected header: {header}")
        for r in rows[1:]:
            if not r or r[0] is None:
                continue
            comps = {
                "adoption": _num(r[8]),
                "trust": _num(r[9]),
                "health": _num(r[10]),
                "ecosystem": _num(r[11]),
                "domain_fit": _num(r[12]),
                "governance": _num(r[13]),
                "docs": _num(r[14]),
            }
            rows_out.append(
                SeedRow(
                    domain=domain,
                    rank=int(r[0]),
                    name=str(r[1]).strip(),
                    category=_cell(r[2]),
                    vendor=_cell(r[3]),
                    ecosystem=_cell(r[4]),
                    url=_cell(r[5]),
                    popularity=_cell(r[6]),
                    license=_cell(r[7]),
                    components=comps,
                    workbook_score=float(r[15]),
                    why=_cell(r[16]),
                    primary_source=_cell(r[17]),
                    last_verified=str(r[18])[:10] if r[18] else None,
                )
            )
    return snapshot_date, rows_out, source_catalog


def import_seed(path: str | Path, tolerance: float = 0.1, cfg: ScoringConfig | None = None) -> SeedImport:
    """Import the workbook and verify every score reproduces (FR-16 acceptance)."""
    cfg = cfg or ScoringConfig.load()
    snapshot_date, rows, source_catalog = read_workbook(path)
    taxonomy = load_config("taxonomy.json")
    valid_domains = set(taxonomy["domains"])

    resources: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    max_err = 0.0

    for row in rows:
        if row.domain not in valid_domains:
            errors.append(f"unknown domain '{row.domain}'")
            continue
        key = f"{slug(row.vendor or '')}:{slug(row.name)}"
        res = resources.get(key)
        if res is None:
            res = new_resource(
                name=row.name,
                source_system=SOURCE_SYSTEM,
                source_url=row.primary_source,
                vendor=row.vendor,
                url=row.url or row.primary_source,
                resource_type=infer_resource_type(row.category),
                category=row.category,
                description=row.why,
                key=key,
                raw={"sheet": row.domain, "rank": row.rank},
            )
            res["deployment"] = row.ecosystem
            res["license"] = row.license
            res["commercial"] = row.license not in (
                None,
                "MIT",
                "Apache-2.0",
                "AGPL-3.0",
                "GPL-3.0",
                "BSD-3-Clause",
            )
            res["signals"] = {"observed_popularity": row.popularity}
            res["components"] = dict(row.components)
            res["seed"] = {"snapshot_date": snapshot_date, "last_verified": row.last_verified}
            resources[key] = res
        else:
            add_source(res, SOURCE_SYSTEM, row.primary_source, raw={"sheet": row.domain, "rank": row.rank})

        result = score_for_domain(row.components, row.domain, cfg)
        err = abs(result.score - row.workbook_score)
        max_err = max(max_err, err)
        if err > tolerance + 1e-9:
            errors.append(
                f"{row.domain} #{row.rank} {row.name}: workbook {row.workbook_score} vs recomputed {result.score} "
                f"(profile {result.profile})"
            )
        res["domains"][row.domain] = {
            "rank": row.rank,
            "score": row.workbook_score,
            "recomputed": result.score,
            "profile": result.profile,
            "score_version": result.version,
            "evidence_coverage": result.evidence_coverage,
            "confidence": 1.0,
            "source": SOURCE_SYSTEM,
        }

    if len(rows) != 1200:
        errors.append(f"expected 1,200 ranked rows, found {len(rows)}")
    for name in COMPONENTS:
        missing = [r for r in rows if r.components.get(name) is None]
        if missing:
            errors.append(f"{len(missing)} rows have no '{name}' component")

    return SeedImport(
        snapshot_date=snapshot_date,
        resources=resources,
        rows=rows,
        max_score_error=round(max_err, 3),
        errors=errors,
        source_catalog=source_catalog,
    )
